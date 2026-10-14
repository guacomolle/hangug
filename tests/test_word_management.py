import io
import unittest
import unicodedata
from pathlib import Path

from flask import Flask
from docx import Document
from openpyxl import Workbook

from models import db, File, Word, ErrorLog, HardWord
from routes import bp


class WordManagementTests(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__, template_folder=str(Path(__file__).resolve().parents[1] / 'templates'))
        self.app.config.update(TESTING=True, SECRET_KEY='test', SQLALCHEMY_DATABASE_URI='sqlite://')
        db.init_app(self.app)
        self.app.register_blueprint(bp)
        self.context = self.app.app_context()
        self.context.push()
        db.create_all()
        first, second = File(name='First'), File(name='Second')
        db.session.add_all([first, second])
        db.session.flush()
        self.file_id, self.other_id = first.id, second.id
        word = Word(file_id=first.id, korean='사과', russian='яблоко')
        db.session.add_all([word, Word(file_id=second.id, korean='배', russian='груша')])
        db.session.commit()
        self.word_id = word.id
        self.client = self.app.test_client()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def test_import_each_format_scopes_and_deduplicates(self):
        rows = [('korean', 'russian'), (' 사과 ', 'другой перевод'),
                (unicodedata.normalize('NFD', '사과'), 'яблоко'),
                ('배', 'груша'), ('배', 'лодка')]
        for extension in ('csv', 'xlsx', 'docx'):
            with self.subTest(extension=extension):
                Word.query.filter_by(file_id=self.file_id, korean='배').delete()
                db.session.commit()
                buffer = io.BytesIO()
                if extension == 'csv':
                    buffer.write('\n'.join(','.join(row) for row in rows).encode('utf-8'))
                elif extension == 'xlsx':
                    workbook = Workbook()
                    for row in rows:
                        workbook.active.append(row)
                    workbook.save(buffer)
                else:
                    document = Document()
                    table = document.add_table(rows=0, cols=2)
                    for row in rows:
                        cells = table.add_row().cells
                        cells[0].text, cells[1].text = row
                    document.save(buffer)
                original = buffer.getvalue()
                response = self.client.post(f'/file/{self.file_id}/add-words',
                    data={'file': (io.BytesIO(original), f'words.{extension}')}, follow_redirects=True)
                self.assertEqual(response.status_code, 200)
                self.assertIn('Добавлено слов: 1. Пропущено повторов: 3', response.get_data(as_text=True))
                self.assertEqual(Word.query.filter_by(file_id=self.file_id).count(), 2)
                self.assertEqual(db.session.get(Word, self.word_id).russian, 'яблоко')
                self.assertEqual(Word.query.filter_by(file_id=self.other_id).count(), 1)
                # Reimporting the same bytes is a no-op.
                self.client.post(f'/file/{self.file_id}/add-words',
                    data={'file': (io.BytesIO(original), f'words.{extension}')})
                self.assertEqual(Word.query.filter_by(file_id=self.file_id).count(), 2)

    def test_edit_and_validation(self):
        url = f'/file/{self.file_id}/word/{self.word_id}/edit'
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertEqual(self.client.post(url, data={'korean': ' 배 ', 'russian': ' груша '}).status_code, 302)
        self.assertEqual(db.session.get(Word, self.word_id).korean, '배')
        for data in ({'korean': ' ', 'russian': 'x'}, {'korean': 'x', 'russian': 'x' * 256}):
            self.assertEqual(self.client.post(url, data=data).status_code, 400)
        self.assertEqual(db.session.get(Word, self.word_id).russian, 'груша')

    def test_delete_cascades_and_ends_affected_training(self):
        db.session.add_all([ErrorLog(word_id=self.word_id), HardWord(word_id=self.word_id)])
        db.session.commit()
        with self.client.session_transaction() as state:
            state.update(mode='memorize', current=self.word_id, queue=[], history=[])
        url = f'/file/{self.file_id}/word/{self.word_id}/delete'
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertEqual(self.client.post(url, follow_redirects=True).status_code, 200)
        self.assertIsNone(db.session.get(Word, self.word_id))
        self.assertEqual(ErrorLog.query.count(), 0)
        self.assertEqual(HardWord.query.count(), 0)
        with self.client.session_transaction() as state:
            self.assertNotIn('mode', state)

    def test_wrong_library_cannot_change_word(self):
        for action in ('edit', 'delete'):
            response = self.client.post(f'/file/{self.other_id}/word/{self.word_id}/{action}',
                                        data={'korean': 'x', 'russian': 'y'})
            self.assertEqual(response.status_code, 404)
        self.assertEqual(db.session.get(Word, self.word_id).korean, '사과')

    def test_clear_errors_requires_confirmation_and_preserves_words(self):
        other_word = Word.query.filter_by(file_id=self.other_id).first()
        db.session.add_all([ErrorLog(word_id=self.word_id), ErrorLog(word_id=other_word.id),
                            HardWord(word_id=self.word_id), HardWord(word_id=other_word.id)])
        db.session.commit()
        self.assertEqual(self.client.get('/errors/clear').status_code, 200)
        self.client.post('/errors/clear')
        self.assertEqual(ErrorLog.query.count(), 2)
        self.assertEqual(HardWord.query.count(), 2)
        with self.client.session_transaction() as state:
            state.update(mode='test', scope='errors', current=self.word_id, queue=[other_word.id])
        response = self.client.post('/errors/clear', data={'confirm': 'clear'}, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(ErrorLog.query.count(), 0)
        self.assertEqual(HardWord.query.count(), 0)
        self.assertEqual(Word.query.count(), 2)
        self.assertEqual(File.query.count(), 2)
        with self.client.session_transaction() as state:
            self.assertNotIn('mode', state)
        from quiz import record_error
        record_error(self.word_id)
        self.assertEqual(ErrorLog.query.count(), 1)

    def test_clear_empty_errors_preserves_regular_training(self):
        with self.client.session_transaction() as state:
            state.update(mode='memorize', scope='file', current=self.word_id)
        response = self.client.post('/errors/clear', data={'confirm': 'clear'}, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        with self.client.session_transaction() as state:
            self.assertEqual(state['current'], self.word_id)


if __name__ == '__main__':
    unittest.main()
