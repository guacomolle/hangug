import csv
import io
from typing import List, Tuple

from docx import Document
from openpyxl import load_workbook

HEADER_KEYWORDS = {
    'korean', 'корейский', 'корейское', 'слово', 'kr',
    'russian', 'русский', 'перевод', 'translation', 'ru',
}


def _is_header_row(a: str, b: str) -> bool:
    a = (a or '').strip().lower()
    b = (b or '').strip().lower()
    return a in HEADER_KEYWORDS or b in HEADER_KEYWORDS


def _clean_rows(rows: List[Tuple[str, str]]) -> List[Tuple[str, str]]:
    if rows and _is_header_row(rows[0][0], rows[0][1]):
        rows = rows[1:]

    result = []
    for kr, ru in rows:
        kr = (kr or '').strip()
        ru = (ru or '').strip()
        if kr and ru:
            result.append((kr, ru))
    return result


def parse_csv(file_stream) -> List[Tuple[str, str]]:
    raw = file_stream.read()
    text = None
    for encoding in ('utf-8-sig', 'cp1251', 'utf-8'):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError('Не удалось определить кодировку CSV-файла')

    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect)
    rows = []
    for row in reader:
        if len(row) < 2:
            continue
        rows.append((row[0], row[1]))
    return _clean_rows(rows)


def parse_xlsx(file_stream) -> List[Tuple[str, str]]:
    wb = load_workbook(file_stream, read_only=True, data_only=True)
    sheet = wb.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        if not row or len(row) < 2:
            continue
        kr, ru = row[0], row[1]
        rows.append((str(kr) if kr is not None else '', str(ru) if ru is not None else ''))
    return _clean_rows(rows)


def parse_docx(file_stream) -> List[Tuple[str, str]]:
    document = Document(file_stream)
    rows = []

    for table in document.tables:
        for tr in table.rows:
            cells = tr.cells
            if len(cells) < 2:
                continue
            rows.append((cells[0].text, cells[1].text))

    if not rows:
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            for sep in ('\t', ' - ', ' – ', '=', ';', ','):
                if sep in text:
                    kr, _, ru = text.partition(sep)
                    rows.append((kr, ru))
                    break

    return _clean_rows(rows)


def parse_file(filename: str, file_stream) -> List[Tuple[str, str]]:
    name = filename.lower()
    if name.endswith('.csv'):
        pairs = parse_csv(file_stream)
    elif name.endswith('.xlsx'):
        pairs = parse_xlsx(file_stream)
    elif name.endswith('.docx'):
        pairs = parse_docx(file_stream)
    else:
        raise ValueError('Неподдерживаемый формат файла')

    if not pairs:
        raise ValueError(
            'Не удалось найти слова в файле. Файл должен содержать два столбца: '
            'корейское слово и перевод на русский.'
        )

    return pairs
